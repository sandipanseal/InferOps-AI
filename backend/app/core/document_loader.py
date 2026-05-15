from pathlib import Path
from tempfile import NamedTemporaryFile

from docx import Document
from fastapi import UploadFile
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def get_file_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def validate_supported_file(filename: str) -> None:
    extension = get_file_extension(filename)

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(
            f"Unsupported file type '{extension}'. Supported file types: {supported}"
        )


async def save_upload_to_temp(upload_file: UploadFile) -> str:
    suffix = get_file_extension(upload_file.filename or "")

    with NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
        content = await upload_file.read()
        temp_file.write(content)
        return temp_file.name


def extract_pdf_text(path: str) -> str:
    reader = PdfReader(path)
    pages: list[str] = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        if text.strip():
            pages.append(f"[Page {page_number}]\n{text.strip()}")

    return "\n\n".join(pages).strip()


def extract_docx_text(path: str) -> str:
    document = Document(path)
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table_number, table in enumerate(document.tables, start=1):
        parts.append(f"\n[Table {table_number}]")

        for row in table.rows:
            row_values = []

            for cell in row.cells:
                value = cell.text.strip()
                if value:
                    row_values.append(value)

            if row_values:
                parts.append(" | ".join(row_values))

    return "\n".join(parts).strip()


def extract_plain_text(path: str) -> str:
    return Path(path).read_text(encoding="utf-8", errors="ignore").strip()


def extract_text_from_file(path: str, filename: str) -> str:
    extension = get_file_extension(filename)

    if extension == ".pdf":
        return extract_pdf_text(path)

    if extension == ".docx":
        return extract_docx_text(path)

    if extension in {".txt", ".md"}:
        return extract_plain_text(path)

    raise ValueError(f"Unsupported file type: {extension}")


async def load_upload_file_text(upload_file: UploadFile) -> dict:
    filename = upload_file.filename or "uploaded_file"
    validate_supported_file(filename)

    temp_path = await save_upload_to_temp(upload_file)

    try:
        text = extract_text_from_file(temp_path, filename)
    finally:
        Path(temp_path).unlink(missing_ok=True)

    if not text.strip():
        raise ValueError(
            "No readable text could be extracted. "
            "This may be a scanned/image-only PDF. OCR support can be added later."
        )

    return {
        "filename": filename,
        "text": text,
        "characters": len(text),
    }